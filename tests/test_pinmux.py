import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from soc_itg.pinmux import read_pinmux, validate


def _make_workbook(path):
    wb = openpyxl.Workbook()
    pads = wb.active
    pads.title = "pads"
    pads.append(["pad", "type", "wrapper", "security", "DS[3]", "PULL[2]"])
    pads.append(["PAD_A", "digital", "IO_BIDIR", "secure", "0b011", "0b10"])
    pads.append(["PAD_B", "digital", "IO_BIDIR", "nonsecure", "0b001", "0b00"])
    pads.append(["PAD_ANA", "analog", "ANA", "nonsecure", None, None])
    fs = wb.create_sheet("function_share")
    fs.append(["group", "pad", "function0", "function1"])
    fs.append(["A", "PAD_A", "cpu0.tx", "cpu0.gpio0"])
    ps = wb.create_sheet("padshare")
    ps.append(["pad", "pin0", "pin1"])
    ps.append(["PAD_B", "cpu0.clk", "cpu0.gpio1"])
    bs = wb.create_sheet("bist_share")
    bs.append(["pad", "bist_func0"])
    bs.append(["PAD_A", "cpu0.bist0"])
    bs.append(["PAD_B", "cpu0.bist1"])
    bs.append(["BIST_ENTRY"])
    bs.append(["bist_mode", "pad", "value"])
    bs.append([0, "PAD_A", 1])
    wb.save(path)


class PinmuxReaderTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.xlsx = Path(self.tmp.name) / "pinmux.xlsx"
        _make_workbook(self.xlsx)

    def tearDown(self):
        self.tmp.cleanup()

    def test_read_pads(self):
        pm = read_pinmux(self.xlsx)
        self.assertEqual(len(pm.pads), 3)
        pad_a = pm.pad("PAD_A")
        self.assertEqual(pad_a.type, "digital")
        self.assertEqual(pad_a.security, "secure")
        self.assertEqual(pad_a.config["DS"].width, 3)
        self.assertEqual(pad_a.config["DS"].default, 3)
        self.assertEqual(pad_a.config["PULL"].default, 2)
        self.assertEqual(pm.pad("PAD_ANA").config, {})

    def test_read_modes(self):
        pm = read_pinmux(self.xlsx)
        self.assertEqual(len(pm.function_share), 1)
        self.assertEqual(pm.function_share[0].group, "A")
        self.assertEqual(pm.function_share[0].functions["function0"], "cpu0.tx")
        self.assertEqual(len(pm.padshare), 1)
        self.assertEqual(pm.padshare[0].pins["pin1"], "cpu0.gpio1")
        self.assertEqual(len(pm.bist_share), 2)
        self.assertEqual(len(pm.bist_entries), 1)
        self.assertEqual(pm.bist_entries[0].pad, "PAD_A")

    def test_validate_ok(self):
        pm = read_pinmux(self.xlsx)
        findings = validate(pm)
        errors = [f for f in findings if f["level"] == "error"]
        self.assertEqual(errors, [])


if __name__ == "__main__":
    unittest.main()
