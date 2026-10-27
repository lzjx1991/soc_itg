import re
from pathlib import Path
import openpyxl
from .model import Pad, ConfigField, FunctionShareEntry, PadShareEntry, BistShareEntry, BistEntry, Pinmux

_CONFIG_RE = re.compile(r"^\s*([A-Za-z_][A-Za-z0-9_]*)\s*\[\s*(\d+)\s*\]\s*$")
_BIST_MARKERS = {"bist_entry", "[bist_entry]"}


def _norm(value):
    if value is None:
        return None
    if isinstance(value, str):
        v = value.strip()
        return v or None
    return value


def _sheet(workbook, name):
    for title in workbook.sheetnames:
        if title.strip().lower() == name.lower():
            return workbook[title]
    return None


def _rows(worksheet):
    return [[_norm(c) for c in row] for row in worksheet.iter_rows(values_only=True)]


def _col(header, name):
    for i, h in enumerate(header):
        if h is not None and str(h).strip().lower() == name.lower():
            return i
    return -1


def _cell(row, index):
    if 0 <= index < len(row):
        return row[index]
    return None


def _bits(value, width):
    if value is None:
        return 0
    if isinstance(value, (int, float)):
        return int(value) & ((1 << width) - 1)
    text = str(value).strip().replace(",", "")
    if not text:
        return 0
    low = text.lower()
    try:
        if low.startswith("0b"):
            return int(text, 2) & ((1 << width) - 1)
        if low.startswith("0x"):
            return int(text, 16) & ((1 << width) - 1)
        if set(text) <= {"0", "1"} and len(text) == width:
            return int(text, 2) & ((1 << width) - 1)
        return int(text, 10) & ((1 << width) - 1)
    except ValueError:
        return 0


def _to_int(value):
    if isinstance(value, (int, float)):
        return int(value)
    return int(str(value).strip(), 0)


def _parse_pads(worksheet):
    rows = _rows(worksheet)
    if not rows:
        return []
    header = rows[0]
    pad_i = _col(header, "pad")
    if pad_i < 0:
        return []
    type_i = _col(header, "type")
    wrapper_i = _col(header, "wrapper")
    security_i = _col(header, "security")
    config_cols = []
    for i, h in enumerate(header):
        if h is None:
            continue
        m = _CONFIG_RE.match(str(h))
        if m:
            config_cols.append((i, m.group(1), int(m.group(2))))
    result = []
    last_pad = None
    for row in rows[1:]:
        pad = _cell(row, pad_i) or last_pad
        if pad:
            last_pad = pad
        if not pad:
            continue
        ptype = str(_cell(row, type_i) or "digital").strip().lower()
        is_digital = ptype != "analog"
        security = str(_cell(row, security_i) or "nonsecure").strip().lower()
        config = {}
        if is_digital:
            for i, name, width in config_cols:
                config[name] = ConfigField(name=name, width=width, default=_bits(_cell(row, i), width))
        result.append(Pad(
            name=str(pad),
            type="digital" if is_digital else "analog",
            wrapper=str(_cell(row, wrapper_i) or ""),
            security="secure" if security == "secure" else "nonsecure",
            config=config,
        ))
    return result


def _parse_function_share(worksheet):
    rows = _rows(worksheet)
    if not rows:
        return []
    header = rows[0]
    group_i = _col(header, "group")
    pad_i = _col(header, "pad")
    func_cols = [i for i, h in enumerate(header) if i not in (group_i, pad_i) and h is not None]
    result = []
    last_group = None
    last_pad = None
    for row in rows[1:]:
        group = _cell(row, group_i) or last_group or "default"
        last_group = group
        pad = _cell(row, pad_i) or last_pad
        if pad:
            last_pad = pad
        if not pad:
            continue
        functions = {str(header[i]): str(_cell(row, i)) for i in func_cols if _cell(row, i)}
        result.append(FunctionShareEntry(group=str(group), pad=str(pad), functions=functions))
    return result


def _parse_padshare(worksheet):
    rows = _rows(worksheet)
    if not rows:
        return []
    header = rows[0]
    pad_i = _col(header, "pad")
    pin_cols = [i for i, h in enumerate(header) if i != pad_i and h is not None]
    result = []
    last_pad = None
    for row in rows[1:]:
        pad = _cell(row, pad_i) or last_pad
        if pad:
            last_pad = pad
        if not pad:
            continue
        pins = {str(header[i]): str(_cell(row, i)) for i in pin_cols if _cell(row, i)}
        result.append(PadShareEntry(pad=str(pad), pins=pins))
    return result


def _parse_bist_table(rows):
    if not rows:
        return []
    header = rows[0]
    pad_i = _col(header, "pad")
    func_cols = [i for i, h in enumerate(header) if i != pad_i and h is not None]
    result = []
    last_pad = None
    for row in rows[1:]:
        pad = _cell(row, pad_i) or last_pad
        if pad:
            last_pad = pad
        if not pad:
            continue
        functions = {str(header[i]): str(_cell(row, i)) for i in func_cols if _cell(row, i)}
        result.append(BistShareEntry(pad=str(pad), functions=functions))
    return result


def _parse_bist_entries(header, rows):
    if not header:
        return []
    mode_i = _col(header, "bist_mode")
    pad_i = _col(header, "pad")
    value_i = _col(header, "value")
    if min(mode_i, pad_i, value_i) < 0:
        return []
    result = []
    for row in rows:
        if _cell(row, mode_i) is None or _cell(row, pad_i) is None:
            continue
        result.append(BistEntry(bist_mode=_to_int(_cell(row, mode_i)), pad=str(_cell(row, pad_i)), value=_to_int(_cell(row, value_i) or 0)))
    return result


def _parse_bist_share(worksheet):
    rows = _rows(worksheet)
    if not rows:
        return [], []
    table_rows = []
    entry_rows = []
    entry_header = None
    state = "table"
    for row in rows:
        first = next((c for c in row if c is not None), None)
        if first is not None and str(first).strip().lower() in _BIST_MARKERS:
            state = "entry_header"
            continue
        if state == "entry_header":
            entry_header = row
            state = "entry"
            continue
        if state == "entry":
            if first is not None:
                entry_rows.append(row)
            continue
        table_rows.append(row)
    return _parse_bist_table(table_rows), _parse_bist_entries(entry_header, entry_rows)


def read_pinmux(path):
    workbook = openpyxl.load_workbook(Path(path), data_only=True)
    pads_ws = _sheet(workbook, "pads")
    fs_ws = _sheet(workbook, "function_share")
    ps_ws = _sheet(workbook, "padshare")
    bs_ws = _sheet(workbook, "bist_share")
    pads = _parse_pads(pads_ws) if pads_ws is not None else []
    function_share = _parse_function_share(fs_ws) if fs_ws is not None else []
    padshare = _parse_padshare(ps_ws) if ps_ws is not None else []
    bist_share, bist_entries = _parse_bist_share(bs_ws) if bs_ws is not None else ([], [])
    return Pinmux(pads=pads, function_share=function_share, padshare=padshare, bist_share=bist_share, bist_entries=bist_entries)
