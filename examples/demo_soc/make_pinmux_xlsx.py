from pathlib import Path
import openpyxl

OUT = Path(__file__).resolve().parent / "pinmux.xlsx"

wb = openpyxl.Workbook()
pads = wb.active
pads.title = "pads"
pads.append(["pad", "type", "wrapper", "security", "DS[3]", "PULL[2]", "SCHMITT[1]"])
pads.append(["PAD_GPIO0", "digital", "IO_BIDIR_WRAP", "secure", "0b011", "0b10", "0"])
pads.append(["PAD_GPIO1", "digital", "IO_BIDIR_WRAP", "nonsecure", "0b001", "0b00", "0"])
pads.append(["PAD_SPI0", "digital", "IO_BIDIR_WRAP", "nonsecure", "0b010", "0b00", "1"])
pads.append(["PAD_TM0", "digital", "IO_IN_WRAP", "nonsecure", "0b000", "0b01", "0"])
pads.append(["PAD_VDDA", "analog", "ANA_PAD_WRAP", "nonsecure", None, None, None])

fs = wb.create_sheet("function_share")
fs.append(["group", "pad", "function0", "function1"])
fs.append(["A", "PAD_GPIO0", "cpu0.paddr", "uart0.prdata"])
fs.append(["A", "PAD_GPIO1", "cpu0.psel", "uart0.pready"])

ps = wb.create_sheet("padshare")
ps.append(["pad", "pin0", "pin1"])
ps.append(["PAD_SPI0", "cpu0.pwrite", "uart0.irq"])
ps.append(["PAD_TM0", "cpu0.penable", "uart0.pwdata"])

bs = wb.create_sheet("bist_share")
bs.append(["pad", "bist_func0", "bist_func1"])
bs.append(["PAD_GPIO0", "cpu0.bist0", "uart0.bist0"])
bs.append(["PAD_GPIO1", "cpu0.bist1", "uart0.bist1"])
bs.append(["PAD_SPI0", "cpu0.bist2", "uart0.bist2"])
bs.append(["PAD_TM0", "cpu0.bist3", "uart0.bist3"])
bs.append(["BIST_ENTRY"])
bs.append(["bist_mode", "pad", "value"])
bs.append([0, "PAD_TM0", 1])

wb.save(OUT)
print(f"wrote {OUT}")
